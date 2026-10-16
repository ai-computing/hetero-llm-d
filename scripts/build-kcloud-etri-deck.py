#!/usr/bin/env python3
"""Build the four-slide ETRI collaboration briefing (editable PPTX + static PDF).

Dependencies: python-pptx, cairosvg, pypdf, Pillow.
The PPTX embeds the original animated GIF. PDF/previews show its first frame.
"""
from pathlib import Path
from io import BytesIO
from xml.sax.saxutils import escape
import base64
import json
import zipfile

import cairosvg
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pypdf import PdfWriter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs/presentations/kcloud-etri-20260908'
OUT.mkdir(parents=True, exist_ok=True)
W, H = 13.333333, 7.5
FONT = 'Noto Sans CJK KR'
C = dict(bg='F5F7FB', ink='14243A', muted='627187', blue='265EDB', teal='007F86',
         light='E8EFFC', line='DCE3ED', white='FFFFFF', navy='13283F', gold='A86A19')
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
prs.core_properties.title = 'K-Cloud | ETRI 이종 가속기 자원관리 연구 기반 및 협업 제안'
prs.core_properties.subject = 'llm-d, lmd-top, 실제 KV 이송 기반 NPU Prefill/Decode 분리'
prs.core_properties.author = 'ETRI'
prs.core_properties.keywords = 'K-Cloud, ETRI, llm-d, lmd-top, NPU, PD disaggregation'
slides = []


class Page:
    def __init__(self, no, label, title, subtitle):
        self.slide = prs.slides.add_slide(prs.slide_layouts[6])
        self.svg = []
        self.no = no
        self.rect(0, 0, W, H, C['bg'])
        self.rect(.48, .36, .09, .20, C['teal'])
        self.text(.69, .30, 10, .30, 'K-CLOUD  /  ETRI     ' + label, 11, C['teal'], True)
        self.text(.55, .92, 12.25, .93, title, 27, C['ink'], True)
        self.text(.57, 1.92, 12.0, .45, subtitle, 13, C['muted'])
        self.rect(.55, 7.10, 12.22, .012, C['line'])
        self.text(.57, 7.18, 11, .20, 'ETRI · K-Cloud 자원관리  |  수행 내용 및 협업 논의  |  2026.09', 8, C['muted'])
        self.text(12.15, 7.15, .6, .25, f'{no:02d} / 04', 9, C['muted'], align='right')
        slides.append(self)

    def rect(self, x, y, w, h, fill, line=None):
        s = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        s.fill.solid()
        s.fill.fore_color.rgb = RGBColor.from_string(fill)
        s.line.fill.background() if line is None else None
        if line:
            s.line.color.rgb = RGBColor.from_string(line)
            s.line.width = Pt(.7)
        self.svg.append(f'<rect x="{x*144}" y="{y*144}" width="{w*144}" height="{h*144}" fill="#{fill}"' + (f' stroke="#{line}"' if line else '') + '/>')
        return s

    def text(self, x, y, w, h, value, size=16, color=None, bold=False, align='left'):
        color = color or C['ink']
        s = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = s.text_frame
        tf.clear()
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        from pptx.enum.text import PP_ALIGN
        for i, line in enumerate(value.split('\n')):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {'left': PP_ALIGN.LEFT, 'right': PP_ALIGN.RIGHT, 'center': PP_ALIGN.CENTER}[align]
            p.space_before = Pt(0)
            p.space_after = Pt(0)
            p.line_spacing = Pt(size * 1.30)
            r = p.add_run()
            r.text = line
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(color)
            rp = r._r.get_or_add_rPr()
            ea = OxmlElement('a:ea')
            ea.set('typeface', FONT)
            rp.append(ea)
            xx = x if align == 'left' else x+w if align == 'right' else x+w/2
            anchor = {'left': 'start', 'right': 'end', 'center': 'middle'}[align]
            yy = y*144 + size*2*1.08 + i*size*2*1.30
            self.svg.append(f'<text x="{xx*144}" y="{yy}" font-family="{FONT}" font-size="{size*2}" font-weight="{700 if bold else 400}" fill="#{color}" text-anchor="{anchor}">{escape(line)}</text>')
        return s

    def picture(self, path, x, y, w, h):
        self.slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
        with Image.open(path) as im:
            frame = im.convert('RGB')
            b = BytesIO()
            frame.save(b, format='PNG')
        encoded = base64.b64encode(b.getvalue()).decode()
        self.svg.append(f'<image x="{x*144}" y="{y*144}" width="{w*144}" height="{h*144}" href="data:image/png;base64,{encoded}"/>')

    def notes(self, content):
        self.note = content
        self.slide.notes_slide.notes_text_frame.text = content

    def export(self):
        svg = '<svg xmlns="http://www.w3.org/2000/svg" width="1920" height="1080" viewBox="0 0 1920 1080">' + ''.join(self.svg) + '</svg>'
        path = OUT / f'slide-{self.no:02d}.svg'
        path.write_text(svg, encoding='utf-8')
        cairosvg.svg2png(bytestring=svg.encode(), write_to=str(OUT / f'slide-{self.no:02d}.png'))
        cairosvg.svg2pdf(bytestring=svg.encode(), write_to=str(OUT / f'slide-{self.no:02d}.pdf'))


p = Page(1, '연구 기반의 전체상', '이종 가속기 자원관리 정책을 검증할\n서빙·관측·실험 기반을 구축했습니다',
         'K-Cloud 관점의 핵심 질문: 어떤 요청을, 어떤 가속기에, 어떤 실행 방식으로 배치할 것인가?')
cards = [
    (.55, '01  서빙 기반', '이종 자원을 하나의\n서빙 체계로 연결', 'Kubernetes + llm-d\nGPU · Rebellions · Furiosa · CPU', '배치·라우팅 정책을 적용할 실행 환경', C['blue']),
    (4.70, '02  관측·운영', '자원에서 요청 경로까지\n상태를 연결해 관측', 'lmd-top + Prometheus\n가속기 · 모델 · EPP · Gateway', '병목을 찾고 정책 효과를 확인할 도구', C['teal']),
    (8.85, '03  실행 방식 실험', 'NPU 간 역할 분리와\n실제 KV 이송 구현', 'RNGD ↔ ATOM Prefill / Decode\n벤치마크 · 런타임 · llm-d 배포', '실행 방식별 비용·효율을 비교할 근거', C['blue']),
]
for x, kicker, title, body, result, accent in cards:
    p.rect(x, 2.63, 3.94, 3.12, C['white'])
    p.rect(x, 2.63, 3.94, .045, accent)
    p.text(x+.22, 2.87, 3.50, .3, kicker, 12, accent, True)
    p.text(x+.22, 3.40, 3.50, .9, title, 21, bold=True)
    p.text(x+.22, 4.49, 3.50, .72, body, 13, C['muted'])
    p.text(x+.22, 5.25, 3.50, .28, result, 11, accent, True)
p.rect(.55, 6.04, 12.24, .65, C['navy'])
p.text(.79, 6.19, 11.8, .34, '다음 연구 연결점     관측 데이터  →  배치·분리 정책  →  실제 서빙 실험  →  SLO·자원 효율 평가', 15, C['white'], True)
p.notes('발표 요지 (약 60초)\nETRI는 K-Cloud 자원관리 연구를 위해 세 가지 기반을 진행했습니다. 첫째는 서로 다른 GPU와 NPU를 연결한 서빙 환경, 둘째는 이를 한 화면에서 관측하고 운영하는 lmd-top, 셋째는 다른 서버에서 수행한 실제 KV 이송 기반 이기종 NPU PD 실험입니다. 이 세 축을 공동연구 정책의 구현·비교에 활용하는 것이 오늘 논의 목적입니다. 세 작업이 완전히 자동화된 폐루프로 통합됐다는 뜻은 아닙니다.\n\n근거: README.md, docs/ENVIRONMENT-GUIDE.md, lmd-top/README.ko.md, lmd-top/CHANGELOG.md, PD disaggregation_HISTORY.md §§1,4,6,12.\nPD 자료는 다른 서버의 작업 이력을 사용자가 제공한 문서로 확인했습니다. 원격 원본 코드나 원시 JSONL을 이 PPT 제작 중 다시 검증하지 않았습니다. 현재 전체 시스템 상시 가동 또는 연구용 플랫폼 제공 확약을 의미하지 않습니다.')

p = Page(2, '성과 ①  서빙·관측·운영', 'llm-d와 lmd-top으로\n이종 서빙의 요청 경로와 자원 상태를 연결했습니다',
         '통합 서빙 환경 구축 + 터미널 관측·운영 도구 개발 + EPP 선택 점수 관측 확장')
gif = ROOT / 'docs/images/lmd-top-ppt.gif'
gx, gy, gw = .55, 2.53, 7.40
gh = gw*1075/1806
p.picture(gif, gx, gy, gw, gh)
for y, n, title, body in [
    (2.58, '01', '모델을 올리고 운영', '공유 모델 스토어 · NPU 컴파일·배포\n스케일 · 재시작 · 롤백'),
    (3.83, '02', '라우팅을 설명할 관측점', 'Gateway → EPP → 모델 → 가속기\n선택 점수 · 큐 · KV 사용률 연결'),
    (5.08, '03', '정책 실험에 연결할 접점', 'Kubernetes 운영 액션 · 상태 JSON\nPrometheus 기반 지연·처리량·자원 지표'),
]:
    p.text(8.29, y, .42, .3, n, 12, C['teal'], True)
    p.text(8.81, y-.03, 3.95, .37, title, 18, bold=True)
    p.text(8.81, y+.46, 3.95, .67, body, 12.5, C['muted'])
p.text(8.29, 6.46, 4.48, .44, '왼쪽 GIF: UI 시연용\n클러스터 구성 기반 · 수치 변화는 합성', 10, C['muted'])
p.notes('발표 요지 (약 70초)\nllm-d 기반으로 가속기별 모델 서버와 모델별 InferencePool/EPP 라우팅을 구성했습니다. lmd-top은 모니터에서 출발해 컴파일·배포와 운영 액션까지 확장했습니다. EPP 확장 플러그인은 실제 엔드포인트별 composite score를 Prometheus로 노출하고, lmd-top에서 해당 값을 연결해서 볼 수 있게 합니다. 연구자가 배치나 라우팅 정책을 제안하면, 어떤 요청 경로에서 어떤 자원 병목이 생기는지 확인할 관측점이 있다는 의미입니다.\n\nGIF: docs/images/lmd-top-ppt.gif, 1806×1075, 25.14초, 무한 반복. 실제 구성 스냅샷을 바탕으로 lmd-top --cast가 수치 변화를 합성했습니다. 실측 성능·실제 부하 영상으로 해석하지 않습니다. PPTX에는 원본 GIF가 포함되고 PDF에는 첫 프레임만 표시됩니다.\n\n근거: lmd-top/README.ko.md, lmd-top/CHANGELOG.md v0.35.0·v0.36.0, lmd-top/contrib/epp-score-observer/README.md (2026-07-08 클러스터 검증 기록), manifests/README-model-store.md, manifests/epp/README.md. 모든 모델이 EPP를 경유하는 것은 아니며, 데이터 표시는 exporter 및 트래픽 조건에 의존합니다. 자동 최적화 컨트롤러가 완성됐다는 의미는 아닙니다.')

p = Page(3, '성과 ②  이기종 NPU 실행 방식', '실제 KV 이송 기반으로\nNPU 간 Prefill / Decode 분리를 구현했습니다',
         'Llama-3.1-8B 기반 실험  |  Prefill: 입력 문맥 처리  /  Decode: 토큰 생성  —  다른 카드·벤더에 배치')
for x, w, title, sub, fill in [
    (.55, 2.00, 'llm-d', 'Gateway · EPP', C['navy']),
    (2.96, 2.54, 'Prefill', 'RNGD 또는 ATOM', C['blue']),
    (5.91, 3.09, '실제 KV cache 이송', 'KVConnector · 공유 메모리', C['teal']),
    (9.41, 3.38, 'Decode', 'RNGD 또는 ATOM', C['blue']),
]:
    p.rect(x, 2.60, w, .93, fill)
    p.text(x+.10, 2.74, w-.20, .31, title, 16, C['white'], True, 'center')
    p.text(x+.08, 3.14, w-.16, .24, sub, 10.5, C['white'], align='center')
for x in [2.59, 5.54, 9.04]:
    p.text(x, 2.89, .32, .35, '→', 19, C['muted'], align='center')

p.rect(.55, 3.83, 5.97, 2.16, C['white'])
p.text(.79, 4.00, 5.40, .32, '구현·최적화 성과', 13, C['teal'], True)
p.text(.79, 4.43, 5.38, .63, '27.00 → 20.10 ms', 29, C['blue'], True)
p.text(.79, 5.13, 5.4, .40, '게이트웨이 경유 TPOT 약 26% 감소', 16, bold=True)
p.text(.79, 5.61, 5.40, .25, '컴파일 설정 개선 전후 · B=4 · 각 20표본 · 8/25 기록', 10, C['muted'])

p.rect(6.77, 3.83, 6.02, 2.16, C['white'])
p.text(7.01, 4.00, 5.50, .32, '자원관리 관점의 발견', 13, C['teal'], True)
p.text(7.01, 4.45, 5.50, .48, 'PD의 전체 처리량 ↑ ≠ 카드당 효율 ↑', 18, bold=True)
p.text(7.01, 5.04, 5.50, .37, 'Solo 1장: 106.6  /  PD 2장: 143.7 tok/s', 15, C['blue'], True)
p.text(7.01, 5.60, 5.50, .25, 'N=8 · 출력 100–160 · 동일 배치 조건 · 후속 최적화 전', 10, C['muted'])
p.rect(.55, 6.22, 12.24, .52, C['light'])
p.text(.79, 6.33, 11.8, .30, '연구 질문     입력·출력 길이와 부하에 따라, 언제 분리하고 어느 가속기에 역할을 맡길 것인가?', 14, C['blue'], True)
p.text(.57, 6.85, 12.1, .18, '근거: PD disaggregation_HISTORY.md §§5.4, 6.5  |  서로 다른 단계의 실험이며 수치를 결합하지 않음', 8, C['muted'])
p.notes('발표 요지 (약 90초)\n다른 서버에서는 공식 엔진 벤치에서 시작해, 자체 런타임으로 KV를 직접 소유하고 실제로 전달하는 단계까지 진행했습니다. 초기의 재프리필 모사와 달리 prefill이 만든 KV를 받아 decode를 이어갑니다. RNGD와 ATOM의 prefill/decode 역할을 교환하는 네 조합, 실제 출력 정합 검증 기록, llm-d 게이트웨이 측정 기록이 있습니다. 슬라이드의 공유 메모리는 현재 서빙 구조이며 별도 TCP 커넥터 검증도 기록돼 있습니다. NIXL·장치 직접 전송 완성을 뜻하지 않습니다.\n\n왼쪽 수치: 제공 문서 §5.4, 42f3e78 (2026-08-25) 커밋의 후속 게이트웨이 측정 기록. B=4, 각 조건 20표본, TPOT 27.00→20.10ms, 감소율 (27.00-20.10)/27.00=25.56%. 해당 절 커널 조건은 S=256·32레이어지만 HTTP 측정의 전체 입력 메타데이터는 별도 확인하지 않아 슬라이드에서 추가하지 않았습니다. TPOT는 배치로 나눈 ms/token과 다릅니다.\n\n오른쪽 수치: 제공 문서 §6.5, 384310c의 동일 배치·closed-loop 조건 재측정 기록. N=8, 출력 100~160, solo 1카드 106.6 tok/s, PD 2카드 143.7 tok/s. 카드당 각각 106.6과 71.85 tok/s이므로 solo가 약 1.48배 높았습니다. 후속 KV·컴파일 최적화 전 결과이고 왼쪽 A/B와 다른 버전·실험입니다. 전체 처리량 이점과 카드당 효율의 trade-off를 보여주는 사례일 뿐 모든 PD가 비효율적이라는 결론이 아닙니다.\n\n근거 문서: PD disaggregation_HISTORY.md §§4,5.4,6,7,11. 이번 제작에서는 다른 서버의 원시 결과를 직접 확인하거나 재측정하지 않았습니다.')

p = Page(4, '공동연구 논의 제안', '구축한 실증 기반 위에서\n배치·KV 관리·SLO 제어 정책을 함께 검증하고자 합니다',
         'ETRI의 구현·계측 기반과 교수진의 알고리즘·모델링 연구를 연결하는 세 가지 접점')
xs = [.55, 3.25, 7.58]
widths = [2.70, 4.33, 5.21]
for x, w, label in zip(xs, widths, ['협업 주제', 'ETRI의 기존 기반', '공동연구 후보 · 제안']):
    p.rect(x, 2.56, w, .47, C['navy'])
    p.text(x+.18, 2.65, w-.36, .27, label, 12, C['white'], True)
rows = [
    ('01  이종 자원 배치', 'RNGD / ATOM 역할 교환 실험\n길이·동시성·전력 측정 하니스', '워크로드별 Solo / PD 선택 정책\n가속기 역할·카드 수의 최적 배분'),
    ('02  KV·배치 관리', '실제 KV 이송·레이어 그룹 스트리밍\nresident KV · 배치 디코드 구현', '전송·계산 중첩과 슬롯별 KV 관리\ncontinuous batching의 병목 개선'),
    ('03  SLO 기반 제어', 'lmd-top · EPP 점수 · 운영 액션\n지연·처리량·자원 관측', 'TTFT·TPOT 목표를 만족하는 제어\n라우팅·스케일링 정책 공동 평가'),
]
for i, (a, b, c) in enumerate(rows):
    y = 3.04 + i*.86
    p.rect(.55, y, 12.24, .84, C['white'] if i%2==0 else 'EDF2F8')
    p.text(.73, y+.24, 2.34, .33, a, 14, C['blue'], True)
    p.text(3.43, y+.13, 3.97, .64, b, 12.5)
    p.text(7.76, y+.13, 4.85, .64, c, 12.5)
p.rect(.55, 5.93, 12.24, .93, C['light'])
p.text(.79, 6.07, 11.70, .30, '첫 공동 실험 제안: 공통 워크로드에서 “언제 Solo / PD를 선택할지” 정책 맵 만들기', 15, C['blue'], True)
p.text(.79, 6.51, 11.70, .23, '조건 통일: 모델·빌드·배치·카드 수   |   평가: TTFT · TPOT · 카드당 tok/s · tok/J', 11.5, C['muted'])
p.notes('발표 요지 (약 75초)\n교수님의 전문 분야에 따라 세 가지 협업 경로를 제안할 수 있습니다. 첫째, 입력·출력 길이와 도착 부하에 따라 Solo/PD와 가속기 역할을 고르는 배치 정책입니다. 둘째, 실제 KV 이송 경로의 중첩과 슬롯별 KV 구조 개선으로 continuous batching의 비용을 줄이는 시스템 연구입니다. 셋째, 관측값과 운영 액션을 연결해 지연 목표를 만족시키는 라우팅·스케일 제어입니다.\n\n우선 공동 실험은 공통 워크로드와 동일 버전의 측정으로 Solo/PD 선택 경계를 만드는 것으로 제안합니다. SLO는 서비스 수준 목표, TTFT는 첫 토큰 지연, TPOT는 토큰 간 생성 지연, tok/J는 에너지 효율입니다. 카드당 효율과 전체 처리량을 함께 비교하고 commit·artifact·batch·scheduler·prefetch를 기록해야 합니다.\n\n이 슬라이드는 협업 제안이며 연구 분담·플랫폼 제공·성능 개선의 확약이 아닙니다. 최신 통합 스윕, 전력 집계 정의 통일 및 현재 운영 상태 재확인이 필요합니다. 원격 문서의 §12는 continuous batching 개선, NIXL backend 및 장치 직접 전송을 미완료로 구분합니다.\n\n근거: PD disaggregation_HISTORY.md §§6.4,6.5,7,11,12; lmd-top/README.ko.md; lmd-top/contrib/epp-score-observer/README.md.')

deck_path = OUT / 'K-Cloud_ETRI_성과와_협업제안_4slides.pptx'
prs.save(deck_path)
pdf = PdfWriter()
for p in slides:
    p.export()
    pdf.append(str(OUT / f'slide-{p.no:02d}.pdf'))
with (OUT / 'K-Cloud_ETRI_성과와_협업제안_4slides.pdf').open('wb') as f:
    pdf.write(f)

notes = '# K-Cloud ETRI 성과·협업 발표 노트\n\n총 4장, 약 5분. 텍스트·도형은 편집 가능하며 2장에 원본 GIF가 삽입돼 있습니다.\nPDF와 미리보기는 같은 레이아웃 데이터로 생성한 정적 참고본입니다. PowerPoint 자체 렌더링 결과는 아닙니다.\n\n'
for p in slides:
    notes += f'## {p.no}장\n\n{p.note}\n\n'
(OUT / '발표노트_및_근거.md').write_text(notes, encoding='utf-8')

# Structural validation: four slides, embedded GIF remains animated, no shapes off-slide.
reopen = Presentation(deck_path)
assert len(reopen.slides) == 4
for si, slide in enumerate(reopen.slides, 1):
    for shape in slide.shapes:
        assert shape.left >= 0 and shape.top >= 0, (si, shape.name)
        assert shape.left+shape.width <= reopen.slide_width+100, (si, shape.name)
        assert shape.top+shape.height <= reopen.slide_height+100, (si, shape.name)
with zipfile.ZipFile(deck_path) as z:
    gifs = [n for n in z.namelist() if n.endswith('.gif')]
    assert len(gifs) == 1
    with Image.open(BytesIO(z.read(gifs[0]))) as im:
        assert im.n_frames > 1
        frames = im.n_frames
print(json.dumps({'pptx': str(deck_path), 'slides': 4, 'embedded_gif_frames': frames,
                  'bytes': deck_path.stat().st_size}, ensure_ascii=False))
