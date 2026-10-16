"""Native, individually editable reconstruction of the approved v10 diagrams."""
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from lxml import etree
import json, zipfile

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs/presentations/ETRI_NPU_Editable_v14'
OUT.mkdir(parents=True,exist_ok=True)
A=ROOT/'docs/presentations/kcloud-etri-research-v7/assets'
LOGOS=ROOT/'docs/presentations/kcloud-etri-visual-v6/assets'
REF=ROOT/'docs/presentations/ETRI_NPU_TopDown_v10/origin_image/slide_01.png'
p=Presentation();p.slide_width=Inches(16);p.slide_height=Inches(9)
BLUE='005EB8'; NAVY='09235D'; MUTED='5C6D84'; PALE='EAF4FD'; GRAY='EDF0F4'; WHITE='FFFFFF'
def txt(s,v,x,y,w,h,size=22,bold=False,c=NAVY,align=None):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));sh.name=v[:45]
    tf=sh.text_frame;tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(v.split('\n')):
        q=tf.paragraphs[0] if not i else tf.add_paragraph();q.text=line
        q.font.name='Noto Sans CJK KR';q.font.size=Pt(size);q.font.bold=bold;q.font.color.rgb=RGBColor.from_string(c)
        if align is not None:q.alignment=align
    return sh
def box(s,x,y,w,h,c=PALE,kind=MSO_SHAPE.ROUNDED_RECTANGLE,line=None):
    sh=s.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(c)
    if line:sh.line.color.rgb=RGBColor.from_string(line);sh.line.width=Pt(2)
    else:sh.line.fill.background()
    # Remove theme-inherited shadows, keeping all geometry native.
    sh._element.spPr.append(etree.Element('{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst'))
    for effect in sh._element.xpath('.//a:effectRef'):effect.set('idx','0')
    if kind==MSO_SHAPE.ROUNDED_RECTANGLE:sh.adjustments[0]=.10
    return sh
def line(s,x,y,x2,y2,c=BLUE,width=2):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x),Inches(y),Inches(x2),Inches(y2));sh.line.color.rgb=RGBColor.from_string(c);sh.line.width=Pt(width)
    for effect in sh._element.xpath('.//a:effectRef'):effect.set('idx','0')
    return sh
def arrow(s,x,y,w=.6,h=.35):return box(s,x,y,w,h,'84B8E8',MSO_SHAPE.RIGHT_ARROW)
def grad(sh,a,b,angle=5400000):
    ns='http://schemas.openxmlformats.org/drawingml/2006/main'
    pr=sh._element.spPr
    for tag in ['solidFill','noFill','gradFill']:
        for old in pr.findall('{'+ns+'}'+tag):pr.remove(old)
    g=etree.Element('{'+ns+'}gradFill',rotWithShape='1');lst=etree.SubElement(g,'{'+ns+'}gsLst')
    for pos,col in [('0',a),('100000',b)]:
        stop=etree.SubElement(lst,'{'+ns+'}gs',pos=pos);etree.SubElement(stop,'{'+ns+'}srgbClr',val=col)
    etree.SubElement(g,'{'+ns+'}lin',ang=str(angle),scaled='1');pr.insert(2,g)
    return sh
def rich(s,kind,x,y,k):
    if kind=='model':
        pts=[(.12,.3),(.12,.75),(.5,.12),(.5,.5),(.5,.88),(.88,.3),(.88,.75)]
        for a,b in [(0,2),(0,3),(1,2),(1,3),(1,4),(2,5),(3,5),(3,6),(4,5),(4,6)]:line(s,x+pts[a][0]*k,y+pts[a][1]*k,x+pts[b][0]*k,y+pts[b][1]*k,'8292AC',2.5)
        for a,b in pts:grad(box(s,x+(a-.11)*k,y+(b-.11)*k,.22*k,.22*k,BLUE,MSO_SHAPE.OVAL),'39A5FF','0062D5')
    elif kind=='chip':
        for q in [.25,.42,.59,.76]:
            for xx,yy,ww,hh in [(x+q*k,y,.06*k,k),(x,y+q*k,k,.06*k)]:box(s,xx,yy,ww,hh,NAVY)
        grad(box(s,x+.11*k,y+.11*k,.82*k,.82*k,BLUE,line=NAVY),'247CE3','063695')
        grad(box(s,x+.2*k,y+.2*k,.64*k,.64*k,BLUE),'419EFF','1260CB')
        center(s,'NPU',x+.2*k,y+.4*k,.64*k,.36*k,15*k,True,WHITE)
    elif kind=='gears':
        for dx,dy,sz in [(0,.37,.62),(.43,0,.55)]:
            box(s,x+dx*k,y+dy*k,sz*k,sz*k,WHITE,MSO_SHAPE.GEAR_6,NAVY)
            box(s,x+(dx+sz*.36)*k,y+(dy+sz*.36)*k,sz*.28*k,sz*.28*k,PALE,MSO_SHAPE.OVAL,NAVY)
    else:
        grad(box(s,x+.12*k,y,.73*k,k,WHITE,MSO_SHAPE.FOLDED_CORNER,NAVY if kind=='code' else '5AA0EE'),'FFFFFF','E9F4FF')
        if kind=='code':center(s,'</>',x+.15*k,y+.31*k,.64*k,.44*k,24*k,True,NAVY)
        elif kind=='search':
            for dy in [.25,.42]:line(s,x+.26*k,y+dy*k,x+.66*k,y+dy*k,NAVY,3)
            box(s,x+.44*k,y+.53*k,.39*k,.39*k,WHITE,MSO_SHAPE.OVAL,NAVY);line(s,x+.77*k,y+.86*k,x+.97*k,y+1.05*k,NAVY,4)
        else:
            box(s,x+.32*k,y+.2*k,.35*k,.35*k,BLUE,MSO_SHAPE.GEAR_6)
            for dy in [.68,.83]:line(s,x+.28*k,y+dy*k,x+.68*k,y+dy*k,'3E78C6',4)
def icon(s,kind,x,y,k=1,c=BLUE):
    # Icons are ordinary editable PowerPoint shapes, never rasterized.
    if kind=='model':
        pts=[(.12,.3),(.12,.75),(.5,.12),(.5,.5),(.5,.88),(.88,.3),(.88,.75)]
        for a,b in [(0,2),(0,3),(1,3),(1,4),(2,5),(3,5),(3,6),(4,6)]:line(s,x+pts[a][0]*k,y+pts[a][1]*k,x+pts[b][0]*k,y+pts[b][1]*k,'A7BAD0',1.4)
        for a,b in pts:box(s,x+(a-.08)*k,y+(b-.08)*k,.16*k,.16*k,c,MSO_SHAPE.OVAL)
    elif kind=='chip':
        box(s,x+.17*k,y+.17*k,.66*k,.66*k,PALE,MSO_SHAPE.ROUNDED_RECTANGLE,c)
        for q in [.3,.5,.7]:
            line(s,x+q*k,y,x+q*k,y+.17*k,c);line(s,x+q*k,y+.83*k,x+q*k,y+k,c)
            line(s,x,y+q*k,x+.17*k,y+q*k,c);line(s,x+.83*k,y+q*k,x+k,y+q*k,c)
    elif kind=='db':box(s,x+.1*k,y,.8*k,k,PALE,MSO_SHAPE.CAN,c)
    elif kind=='bars':
        for i,h in enumerate([.35,.6,.9]):box(s,x+i*.3*k,y+(1-h)*k,.2*k,h*k,c,MSO_SHAPE.RECTANGLE)
    else:
        box(s,x+.12*k,y,.72*k,k,WHITE,MSO_SHAPE.FOLDED_CORNER,c)
        if kind=='code':txt(s,'</>',x+.18*k,y+.32*k,.6*k,.4*k,18*k,True,c,PP_ALIGN.CENTER)
        else:
            for dy in [.3,.5,.7]:line(s,x+.27*k,y+dy*k,x+.65*k,y+dy*k,c,2)
def pic(s,path,x,y,w,h,crop=None):
    iw,ih=Image.open(path).size;l,t,r,b=crop or (0,0,iw,ih);f=min(w/(r-l),h/(b-t));ww,hh=(r-l)*f,(b-t)*f
    sh=s.shapes.add_picture(str(path),Inches(x+(w-ww)/2),Inches(y+(h-hh)/2),width=Inches(ww),height=Inches(hh))
    sh.crop_left=l/iw;sh.crop_top=t/ih;sh.crop_right=(iw-r)/iw;sh.crop_bottom=(ih-b)/ih
    return sh
def start(title,sub=''):
    s=p.slides.add_slide(p.slide_layouts[6]);txt(s,title,.5,.25,13.1,.9,43,True)
    pic(s,REF,13.8,.3,1.7,.65,(1443,30,1625,101))
    if sub:txt(s,sub,.53,1.14,13,.5,21,False,MUTED)
    return s
def center(s,v,x,y,w,h,size=24,bold=True,c=NAVY):return txt(s,v,x,y,w,h,size,bold,c,PP_ALIGN.CENTER)

s=start('이종 자원 관리, 우리는 어디를 개발했나','모델 운영에서 저수준 실행까지')
for y,label in [(1.9,'모델·서비스'),(3.26,'모델 운영'),(4.85,'컴파일·런타임'),(7.18,'GPU·NPU')]:txt(s,label,.55,y,2.2,.5,24,True)
box(s,3,1.85,9.45,1.02,GRAY);icon(s,'file',3.5,2.03,.58);center(s,'사용자 요청 → LLM 추론 서비스',4.5,2.13,7.2,.45,25)
box(s,3,3.13,9.45,1.05,BLUE);center(s,'lmd-top',3.5,3.25,8.4,.43,30,True,WHITE);center(s,'모델 준비 · 배포 · 관측',3.5,3.75,8.4,.35,20,False,WHITE)
for x,k,t in [(3,'file','저수준 실행 분석'),(7.84,'code','커스텀 실행 · KV 전송')]:box(s,x,4.46,4.61,1.02);icon(s,k,x+.2,4.68,.57);txt(s,t,x+.97,4.75,3.55,.45,22,True)
box(s,3,5.69,9.45,.58,GRAY);center(s,'벤더 컴파일러 · 실행 엔진',3.3,5.8,8.8,.35,22)
for y in [2.91,4.21,5.5,6.29]:box(s,7.53,y,.38,.2,'A7BAD0',MSO_SHAPE.UP_ARROW)
box(s,3,6.5,9.45,1.94,GRAY)
for x,name,crop in [(3.2,'NVIDIA GPU',(377,744,647,848)),(6.36,'Furiosa RNGD',(718,737,1002,849)),(9.51,'Rebellions ATOM',(1101,704,1194,854))]:pic(s,REF,x,6.61,2.8,1.25,crop);center(s,name,x,7.97,2.7,.34,19)
box(s,13,1.85,2.5,4.43,GRAY);center(s,'활용 기반',13.2,2.03,2.1,.4,23)
for y,file,title,sub in [(2.73,'kubernetes','Kubernetes','배포·실행 관리'),(4.48,'prometheus','Prometheus','지표 수집')]:pic(s,LOGOS/(file+'.png'),13.73,y,.94,.94);center(s,title,13.1,y+1,2.3,.35,18);center(s,sub,13.1,y+1.4,2.3,.3,14,False)
box(s,13,6.65,.28,.28,PALE);txt(s,'ETRI 분석·개발',13.43,6.6,2.1,.4,17,True,BLUE)
box(s,13,7.26,.28,.28,GRAY);txt(s,'기존 인프라',13.43,7.2,2,.4,17)
txt(s,'작업 범위의 개념도 · 단일 자동 제어기 통합을 의미하지 않음',3,8.64,12,.25,12,False,MUTED)

s=start('장비뿐 아니라 실행 조건도 선택해야 합니다','전체 환경 → 실행 조건')
box(s,.7,2.25,3.3,3.7);icon(s,'model',1.7,2.72,1.3);center(s,'모델',1,4.65,2.7,.5,30)
box(s,5.1,1.93,5.25,4.36);center(s,'조건별 아티팩트',5.35,2.14,4.7,.45,27)
for y,t in [(2.93,'입력 길이 설정'),(4,'배치 설정'),(5.07,'병렬성 설정')]:box(s,5.4,y,4.65,.9,WHITE);icon(s,'file',5.66,y+.15,.52);txt(s,t,6.53,y+.23,3.2,.4,23,True)
arrow(s,4.24,3.82);arrow(s,10.57,3.82)
box(s,11.46,2.25,3.7,3.7);pic(s,REF,12.12,2.66,2.35,2.25,(1101,704,1194,854));center(s,'NPU 실행',11.73,5.22,3.1,.45,28)
for x,k,title,desc in [(.7,'bars','실행 성능','같은 모델도 조건에 따라 달라짐'),(8.1,'file','컴파일 시간 · 파일 크기','실행 성능과 준비 비용을 함께 분석')]:box(s,x,6.8,7.05,1.42);icon(s,k,x+.25,7.09,.6);txt(s,title,x+1.12,7.02,5.6,.45,25,True);txt(s,desc,x+1.12,7.68,5.6,.35,17,False,MUTED)

s=start('모델이 NPU에서 실행되기까지, 무엇을 분석했나')
grad(box(s,5.3,1.26,10.18,5.9),'EAF6FF','DFF0FE');center(s,'우리가 분석한 범위',8.65,1.47,3.55,.48,29,True,BLUE)
line(s,7.48,1.78,8.85,1.78);line(s,12.1,1.78,13.36,1.78);line(s,7.48,1.78,7.48,2);line(s,13.36,1.78,13.36,2)
for x,title,k,desc in [(.55,'모델','model','어떤 연산으로 구성되나?'),(5.94,'실행 준비','file','장비가 실행할 형태로 변환'),(11.33,'NPU 실행','chip','연산과 데이터 이동')]:
    grad(box(s,x,2.16,4.0,3.48),'FFFFFF','F1F6FA');center(s,title,x+.15,2.4,3.7,.55,33)
    rich(s,k,x+1.02,3.03,1.88);center(s,desc,x+.12,5.15,3.76,.44,20,False,MUTED)
grad(arrow(s,4.61,3.54,1.18,.73),'D9EDFF','559BE7',0);grad(arrow(s,10.06,3.54,1.17,.73),'D9EDFF','559BE7',0)
box(s,5.58,5.74,9.64,1.21,'F5FAFE');txt(s,'저수준 IR:',5.86,5.98,4.4,.43,25,True)
txt(s,'연산·메모리 동작을 표현한 중간 단계',5.86,6.46,4.55,.4,19,False,MUTED);line(s,10.52,6.03,10.52,6.7,'C0D0E3',1)
txt(s,'무엇을 계산하고,',11.08,5.98,3.85,.43,25,True);txt(s,'데이터를 어디로 옮기는지 확인',11.08,6.46,3.9,.4,18,False,MUTED)
grad(box(s,2.44,7.66,12.82,1.04,BLUE),'208CEB','005CC1',0)
box(s,3.24,7.85,.44,.44,WHITE,MSO_SHAPE.OVAL);line(s,3.37,8.29,3.57,8.29,WHITE,4)
for xx,yy,xx2,yy2 in [(3.46,7.7,3.46,7.8),(3.13,7.83,3.2,7.9),(3.75,7.83,3.68,7.9),(3.07,8.09,3.18,8.09),(3.76,8.09,3.88,8.09)]:line(s,xx,yy,xx2,yy2,WHITE,2)
line(s,3.37,8.36,3.57,8.36,WHITE,3)
center(s,'실행 비용과 변경 가능한 지점을 파악',4.25,7.98,10.56,.53,30,True,WHITE)

s=start('기존 엔진은 유지하고, 필요한 실행만 바꿉니다')
rich(s,'code',7.35,1.36,.99);txt(s,'우리가 개발한 연결 도구',8.64,1.71,6,.5,29,True,BLUE)
grad(box(s,1.46,2.82,13.1,2.68,GRAY,line='ADBDCF'),'F2F6FA','E5ECF3')
for x,title,c in [(1.72,'기존 실행','DCE3EC'),(6.13,'직접 작성한 코드',BLUE),(10.7,'기존 실행','DCE3EC')]:
    grad(box(s,x,3.22,3.57,1.47,c),'229BFF' if c==BLUE else 'E3EAF1','0063DB' if c==BLUE else 'D9E1EB')
    center(s,title,x+.05,3.74,3.47,.55,31,True,WHITE if c==BLUE else NAVY)
for x in [5.31,9.88]:grad(arrow(s,x,3.8,.74,.38),'E0E6ED','8B9BB1',0)
center(s,'기존 NPU 실행 엔진',3,4.96,10,.48,27)
grad(box(s,7.71,2.4,.48,.77,'A7BAD0',MSO_SHAPE.DOWN_ARROW),'D9E0E8','7F8EA4')
for x,k,v in [(1.24,'gears','엔진 전체를 바꾸지 않고\n필요한 부분만 수정'),(8.14,'search','수정한 코드가 실제\n호출되는 경로 확인')]:
    grad(box(s,x,5.85,6.6,1.56),'E9F5FF','DCEFFF');rich(s,k,x+.77,6.12,1.07);txt(s,v,x+2.19,6.19,4.22,1.07,26,True)
grad(box(s,7.47,7.56,1.08,.32,'A7BAD0',MSO_SHAPE.DOWN_ARROW),'E1E6EC','8493A6')
center(s,'일부 코드 교체에서, NPU 간 실행 분리로',1,8.12,14,.6,28,True,BLUE)

s=start('KV를 전달해, NPU 사이에서 실행을 나눴습니다','입력 처리와 답변 생성을 분리하는 PD 실행')
for x,step,title,desc in [(.55,'01','입력 처리 · Prefill','입력 문맥을 처리하고 KV 생성'),(5.85,'02','KV 전달','중간 상태를 다음 NPU로 전달'),(11.15,'03','답변 생성 · Decode','받은 KV로 다음 토큰 생성')]:
    grad(box(s,x,2.03,4.3,4.85),'F3F9FF','E8F3FD')
    txt(s,step,x+.24,2.22,.75,.42,20,True,BLUE)
    center(s,title,x+.15,2.81,4,.5,27)
    center(s,desc,x+.1,3.49,4.1,.44,18,False,MUTED)
pic(s,REF,1.07,4.2,3.28,1.56,(718,737,1002,849))
grad(box(s,7.23,4.23,1.6,1.58,BLUE,MSO_SHAPE.CAN),'419CFA','0060C8')
center(s,'KV',7.36,4.82,1.34,.5,29,True,WHITE)
pic(s,REF,12.29,4.06,2.05,1.98,(1101,704,1194,854))
for x in [4.58,9.88]:
    grad(arrow(s,x,4.72,1.5,.6),'DAEBFA','4E99E8',0)
    center(s,'KV 전송',x-.05,4.21,1.6,.33,16,True,BLUE)
center(s,'Furiosa RNGD',.8,6.22,3.8,.4,23)
center(s,'직접 구현한 전송 경로',6.04,6.22,3.91,.4,23,True,BLUE)
center(s,'Rebellions ATOM',11.4,6.22,3.8,.4,23)
box(s,.55,7.27,7.35,.99,'EFF5FB');txt(s,'분석한 조건',.86,7.46,1.85,.35,18,True,BLUE);txt(s,'KV 유지 · 연산 배치 · 컴파일',2.84,7.46,4.77,.4,21,True)
box(s,8.2,7.27,7.25,.99,'EFF5FB');txt(s,'실행 조합',8.52,7.46,1.72,.35,18,True,BLUE);txt(s,'동종·이종 NPU / 역할 교환',10.39,7.46,4.8,.4,21,True)

CAP=OUT/'assets'
for title,sub,points,asset in [
    ('lmd-top ① 모델·산출물 관리','모델을 찾고, 실행 가능한 빌드를 구분',
     [('보유 모델 탐색','원본 모델·컴파일 빌드\n선택 항목을 이동'),('Model Zoo 탐색','벤더별 지원 모델과\n준비 상태를 확인'),('Library ↔ Zoo','두 목록을 오가며\n모델을 비교')],CAP/'library-zoo.gif'),
    ('lmd-top ② 컴파일·배포 작업','준비한 모델을 실제 서빙으로 연결',
     [('컴파일 조건 선택','병렬성·길이·배치\n선택 항목과 값 이동'),('배포 조건 선택','복제 수·장비·라우팅\n옵션별 설정 변경'),('적용 전 점검','예상 자원 수요 확인\n적용 없이 취소')],CAP/'options.gif'),
    ('lmd-top ③ 실행·자원 관측','배포 이후 모델과 장비 상태를 함께 확인',
     [('클러스터 상태 요약','모델 서빙과 가속기 현황을\n한 화면에서 확인'),('장비별 자원 관측','사용률·메모리·전력 등\n수집 가능한 지표 표시'),('기존 관측 기반 활용','Kubernetes·Prometheus 연계\n실제 화면 전환을 GIF로 표시')],CAP/'observe.gif')]:
    s=start(title,sub)
    dynamic='①' in title or '②' in title
    for idx,(heading,desc) in enumerate(points):
        y=2.1+idx*1.87
        if dynamic:
            box(s,.5,y,2.82,1.58)
            txt(s,heading,.71,y+.2,2.4,.42,19,True,BLUE)
            txt(s,desc,.71,y+.8,2.4,.67,14,False,MUTED)
        else:
            box(s,.5,y,4.07,1.58);icon(s,['file','code','bars'][idx],.72,y+.2,.42)
            txt(s,heading,1.27,y+.18,3.1,.43,21,True,BLUE);txt(s,desc,.76,y+.77,3.55,.72,17,False,MUTED)
    if dynamic:
        im=pic(s,asset,3.62,1.72,11.9,6.85)
        assert not any((im.crop_left,im.crop_right,im.crop_top,im.crop_bottom))
    else:im=pic(s,asset,4.88,1.86,10.62,6.74)
    assert im.image.blob==asset.read_bytes()
    txt(s,'실제 UI 녹화 · 옵션 이동과 취소만 수행 · 컴파일/배포 미실행',4.98,8.68,10.6,.25,11,False,MUTED)

s=start('다음은 조건·단계·장비를 함께 고르는 정책입니다')
txt(s,'확보한 기반',.65,1.88,8,.5,27,True,BLUE)
for y,k,title,desc in [(2.64,'file','실행 조건','모델 특성 · 실행 환경 분석'),(4.34,'code','실행 단계','저수준 경로 분석 · 커스텀 실행'),(6.04,'chip','모델 운영','모델 준비 · 배포 · 관측')]:box(s,.6,y,8.7,1.38);icon(s,k,.98,y+.27,.75);txt(s,title,2.1,y+.18,6.7,.43,27,True);txt(s,desc,2.1,y+.84,6.7,.37,21,False,MUTED)
arrow(s,9.57,4.63,.65,.45);box(s,10.48,2.04,5,5.63,WHITE,line='7FABD6');center(s,'협업할 질문',10.7,2.29,4.55,.5,28,True,BLUE)
for y,t in [(3.42,'언제 실행을\n나눌까?'),(5.42,'어떤 설정·장비를\n고를까?')]:box(s,10.85,y,4.25,1.55);center(s,t,11,y+.26,3.95,1.1,27)
center(s,'비용 예측과 자원 배치 정책으로 연결',1,8.12,14,.6,29,True,BLUE)

from PIL import ImageFont
# Keep only explicit line breaks on the revised slides. Measure with the actual
# installed font, reserving 8% horizontal room for other PowerPoint renderers.
fit_report=[]
for number in [3,4,5]:
    for sh in p.slides[number-1].shapes:
        if not sh.has_text_frame or not sh.text:continue
        sh.text_frame.word_wrap=False
        for q in sh.text_frame.paragraphs:
            if not q.text or not q.font.size:continue
            fontfile='/usr/share/fonts/opentype/noto/NotoSansCJK-'+('Bold' if q.font.bold else 'Regular')+'.ttc'
            f=ImageFont.truetype(fontfile,1000)
            available=sh.width/12700*.92
            size=q.font.size.pt
            measured=f.getlength(q.text)/1000*size
            if measured>available:
                q.font.size=Pt(size*available/measured)
                fit_report.append((number,q.text,round(size,1),round(q.font.size.pt,1)))
out=OUT/'ETRI_NPU_Editable_v14.pptx';p.save(out)
print('Text fit adjustments:',fit_report)
report=[]
for i,s in enumerate(p.slides,1):report.append({'slide':i,'editable_textboxes':sum(sh.has_text_frame and bool(sh.text) for sh in s.shapes),'objects':len(s.shapes)})
with zipfile.ZipFile(out) as z:
    gifs=[n for n in z.namelist() if n.endswith('.gif')]
    assert len(gifs)==3
    for file in ['library-zoo.gif','options.gif','observe.gif']:assert (CAP/file).read_bytes() in [z.read(g) for g in gifs]
print(json.dumps(report,ensure_ascii=False));print(out)
