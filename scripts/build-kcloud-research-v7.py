from pathlib import Path
from PIL import Image
from pptx.opc.packuri import PackURI
src=Path(__file__).with_name('build-kcloud-visual-v5.py').read_text().replace('kcloud-etri-visual-v5','kcloud-etri-research-v7').replace("p.save(OUT/'ETRI_NPU_모델자원관리_5slides_v5.pptx')",'# save below')
exec(compile(src,__file__,'exec'))
for i,s in enumerate(p.slides,1):s.part._partname=PackURI(f'/ppt/slides/slide{i}.xml')
for _ in range(2):p.slides.add_slide(p.slide_layouts[0])
for s in p.slides:
    for sh in list(s.shapes):sh._element.getparent().remove(sh._element)
def picture(s,name,x,y,w,h=None,crop=None):
    path=assets/name;iw,ih=Image.open(path).size
    l,t,r,b=crop or (0,0,iw,ih)
    if h is not None:
        fit=min(w/(r-l),h/(b-t));nw=(r-l)*fit;nh=(b-t)*fit
        x+=(w-nw)/2;y+=(h-nh)/2;w=nw;h=nh
    sh=s.shapes.add_picture(str(path),Inches(x),Inches(y),width=Inches(w),height=Inches(h or w*(b-t)/(r-l)))
    sh.crop_left=l/iw;sh.crop_right=(iw-r)/iw;sh.crop_top=t/ih;sh.crop_bottom=(ih-b)/ih
    return sh
def our(s,label,x,y,w=3):box(s,x,y,w,.36,PALE);txt(s,label,x+.1,y+.03,w-.2,.28,12,True,PURPLE)
def item(s,ic,title,body,x,y,w=3.7):
    icon(s,ic,x,y,.77);txt(s,title,x,y+1.05,w,.5,24,True);txt(s,body,x,y+1.76,w,1.2,18,False,MUTED)
def nt(s,t):notes(s,t+'\n공통 근거: PD disaggregation_HISTORY.md (별도 개발 서버 작업 기록), lmd-top/README.ko.md. 기록 기반 요약이며 원시 실험 재검증이나 통합 제어기 완성을 주장하지 않는다.')

s=p.slides[0];base(s,1,'ETRI / our work','이종 NPU의 실행 내부부터\n모델·자원 관리까지 개발했습니다')
our(s,'ETRI 분석·구현',.58,2.43,2.4)
for x,ic,title,body in [(.58,'LuWorkflow','저수준 실행 분석','DSL 계산 그래프 구성\n컴파일 조건·메모리 경로 조사'),(4.87,'LuNetwork','실행 분리 런타임','직접 KV를 소유·전달\n동종·이종 NPU의 P/D 연결'),(9.05,'LuActivity','모델·자원 관리','lmd-top 개발\n모델 준비·배포·관측·운영')]:item(s,ic,title,body,x,3.08)
txt(s,'공통 질문: 모델과 실행 단계를 어떤 자원에, 어떤 조건으로 배치할 것인가?',.58,6.52,12.1,.42,19,True,PURPLE)
nt(s,'세 작업축의 공통 맥락을 먼저 제시. 아래의 Kubernetes/Prometheus/벤더 하드웨어를 자체 개발했다고 주장하지 않는다. PD 런타임과 lmd-top은 별도 구현이며 정책 폐루프로 완전 통합한 것이 아니다.')

s=p.slides[1];base(s,2,'Built environment / reused technologies','GPU·Furiosa·Rebellions를 연결한\n실험·서빙 환경을 구성했습니다')
picture(s,'spark.jpg',.58,2.54,3.72,2.04,(1260,690,1710,955))
picture(s,'rngd.png',4.81,2.54,3.72,2.04)
picture(s,'atom.png',9.04,2.54,3.72,2.04)
for x,title,sub in [(.58,'NVIDIA GPU','GB10 기반 DGX Spark'),(4.81,'Furiosa NPU','RNGD'),(9.04,'Rebellions NPU','ATOM')]:
    txt(s,title,x,4.83,3.8,.5,25,True);txt(s,sub,x,5.43,3.8,.39,17,False,MUTED)
for x,n,label in [(.65,'kubernetes','Kubernetes  |  배포·실행 관리'),(6.6,'prometheus','Prometheus  |  지표 수집')]:
    path=ROOT/'docs/presentations/kcloud-etri-visual-v6/assets'/f'{n}.png'
    s.shapes.add_picture(str(path),Inches(x),Inches(6.11),width=Inches(.52));txt(s,label,x+.76,6.2,5.2,.39,19,True)
txt(s,'제조사 공개 제품 이미지 · 보유 장비 실물 사진 아님 · 카드 외형은 리비전에 따라 다를 수 있음',.58,6.82,12,.19,9,False,MUTED)
notes(s,'대표 제품 이미지이며 설치 현장 사진 아님. NVIDIA GB10은 GPU 기반 시스템으로 DGX Spark를 표시. Furiosa RNGD와 Rebellions ATOM은 제조사 공개 제품 이미지이며 보유 카드의 정확한 보드 리비전 일치까지 확인하지 않음. 공식 출처: https://nvidianews.nvidia.com/news/nvidia-announces-dgx-spark-and-dgx-station-personal-ai-computers ; https://furiosa.ai/rngd ; https://rebellions.ai/wp-content/uploads/2025/05/atom.png (공식 media API의 ATOM 이미지). K8s/Prometheus 로고: CNCF artwork. 환경 근거 docs/ENVIRONMENT-GUIDE.md. 논리 구성이지 현재 전 모델 가동을 뜻하지 않는다.')

s=p.slides[2];base(s,3,'ETRI / low-level graph & runtime','저수준 계산 그래프를 구성하고\nKV가 드나드는 실행 경로를 만들었습니다')
for x,title,sub in [(.58,'모델 연산','Attention · MLP · lm_head'),(4.87,'DSL 계산 그래프','레이어 그룹 · shape · KV 입출력'),(9.05,'자체 런타임','가중치 로드 · 그래프 연결')]:
    box(s,x,2.71,3.7,1.63,PALE);txt(s,title,x+.22,3.00,3.26,.5,25,True);txt(s,sub,x+.22,3.67,3.26,.38,16,False,MUTED)
arrow(s,4.35,3.23);arrow(s,8.6,3.23)
icon(s,'LuDatabase',.66,4.99,.8);txt(s,'KV를 명시적인 입출력 텐서로 관리',1.79,4.93,10.7,.52,27,True)
txt(s,'기존 KV를 받는 prefill  ·  KV를 유지하는 decode  ·  CPU 기준 출력 정합 확인',1.79,5.69,10.7,.7,18,False,MUTED)
txt(s,'분석 범위: IR 관련 저수준 표현 중 DSL·계산 그래프·실행 경로',.58,6.61,12,.35,15,False,PURPLE)
nt(s,'§4.1: furiosa-torch/DSL 기반 계산 그래프와 가중치 로딩 엔진. chunked prefill, KV-free full prefill, single-stream resident 및 batched decode, 그룹 그래프 연결. §5.2 config 기반 모델 연산 생성 및 CPU fp32 기준 검증. 원문에 LLVM/MLIR 패스 분석·벤더 내부 IR 변환 패스 직접 수정 근거는 없어 주장하지 않음. PD disaggregation은 실행 분리 방식이지 IR 자체가 아님. 당시 공식 엔진의 KV export/import 훅 제약을 해결하려고 자체 경로 구성.')

s=p.slides[3];base(s,4,'ETRI / PD disaggregation','입력 처리와 답변 생성을 나누고\n실제 KV를 전달해 실행을 이어갔습니다')
for x,title,sub in [(.58,'Prefill worker','입력 문맥 처리'),(9.05,'Decode worker','전달받은 KV로 답변 생성')]:
    box(s,x,2.85,3.7,2.3,PALE);icon(s,'LuCpu',x+.26,3.12,.71);txt(s,title,x+.25,4.03,3.2,.47,24,True);txt(s,sub,x+.25,4.62,3.2,.4,17,False,MUTED)
arrow(s,4.39,3.73);arrow(s,8.49,3.73);icon(s,'LuDatabase',6.12,3.03,1.0)
txt(s,'KVConnector',5.17,4.12,3.25,.49,26,True,PURPLE);txt(s,'KV 소유 · 저장 · 전달 · 로드',4.93,4.75,3.8,.38,16,False,MUTED)
txt(s,'Furiosa ↔ Rebellions',.58,5.7,6.5,.49,28,True)
txt(s,'동종·이종 조합  /  Prefill·Decode 역할 교환',.58,6.34,7.8,.39,18,False,MUTED)
our(s,'직접 구현한 런타임·전송 경로',9.05,5.78,3.7);txt(s,'공유 메모리 · TCP 전송 검증',9.05,6.36,3.9,.4,17,False,MUTED)
nt(s,'§4.2 RngdKVConnector는 vLLM KVConnectorBase_V1 계약 구현. inproc/shm/TCP 전송과 bf16 wire. NIXL은 스텁. §6 HTTProuter→Pworker→공유 /dev/shm→Dworker 배포, Furiosa RNGD/Rebellions ATOM 네 조합. TCP는 별도 프로세스 검증 기록이며 그림이 모두 서버 간 전송을 뜻하지 않음. 초기 re-prefill 에뮬레이션과 실제 KV 전달 PD를 구분. lmd-top에서 자동 제어하는 완성 기능이 아님.')

s=p.slides[4];base(s,5,'ETRI / runtime optimization','연산·메모리·컴파일 경로를 바꿔\n실행 비용이 어디서 생기는지 확인했습니다')
for x,ic,title,body in [(.58,'LuMemoryStick','KV를 장치에 유지','매 토큰 전체 재업로드 축소\n레이어 그룹 스트리밍·프리패치'),(4.87,'LuCpu','연산을 NPU에 배치','lm_head의 호스트 실행 축소\n배치 decode와 데이터 복사 개선'),(9.05,'LuSlidersHorizontal','컴파일 조건 탐색','레이어 그룹·shape 제한 비교\n사전 컴파일·캐시 키 정렬')]:item(s,ic,title,body,x,2.82)
txt(s,'“더 크게 묶기”가 항상 빠르지는 않았습니다 — 실행 조건과 정확도를 함께 확인',.58,6.46,12.1,.52,19,True,PURPLE)
nt(s,'§5.1 resident KV, §5.3 batched decode/lm_head 이동, §5.4 max_num_unique_shapes 및 group 탐색, prebake, cache key; §7 CPU bridge/copy 제거, layer streaming, decode prefetch. 개선 시도와 검증 사실 중심으로 수치 생략. continuous batching도 구현했지만 당시 static보다 느려 기본값 유지. 모든 설정에서 개선됐다는 주장은 피함. 정확도는 기록의 CPU fp32 비교와 출력 일치 범위에 한정.')

s=p.slides[5];base(s,6,'ETRI / lmd-top','모델 준비·배포·관측을 연결하는\n운영 도구 lmd-top을 개발했습니다')
picture(s,'lmd-top-real.png',5.06,2.34,7.0)
for y,title,body in [(2.62,'준비·컴파일','원본과 빌드 산출물 구분\n벤더 컴파일 Job 관리'),(4.01,'배포·작업 이력','타깃 자원 선택\n진행·실패·로그 확인'),(5.4,'관측·운영','모델·요청·장비 상태 연결\n스케일·재시작·정지')]:
    txt(s,title,.58,y,3.5,.46,24,True);txt(s,body,.58,y+.59,3.6,.74,17,False,MUTED)
txt(s,'실제 observe 모드 터미널 녹화 프레임 · 색상 유지 · 데이터 합성 없음',5.06,6.87,7.6,.2,9,False,MUTED)
notes(s,'원래 TUI를 PTY에서 observe 모드로 실행하여 ANSI 색상 포함 asciicast 녹화 후 agg로 렌더링한 실제 프레임. NO_COLOR 환경 변수 제거, soft 테마. --render 문자 재구성이 아니며 OS 데스크톱 캡처도 아님. 메뉴 이동만 했고 compile/deploy를 실행하지 않음. 보이는 Running/Failed 등은 다른 작업의 실제 순간 상태. lmd-top은 벤더 컴파일러 Job을 관리하는 도구이지 자체 컴파일러가 아님. PD 런타임 자동 통합 제어를 주장하지 않음. 근거 lmd-top README.ko.md, assets/lmd-top-real.cast.')

s=p.slides[6];base(s,7,'Collaboration / connect the layers','우리가 만든 실행·관리 기반 위에서\n자원 배치 정책을 함께 연구할 수 있습니다')
for x,title,sub in [(.58,'분석 기반','그래프·컴파일 조건\n연산·데이터 이동 비용'),(4.87,'실행 기반','한 카드 실행 / P·D 분리\n동종·이종 가속기 조합'),(9.05,'관리 기반','모델 준비·배포\n실행·자원 상태 관측')]:
    box(s,x,2.77,3.7,1.99,PALE);txt(s,title,x+.23,3.08,3.24,.5,26,True);txt(s,sub,x+.23,3.81,3.24,.82,19,False,MUTED)
txt(s,'협업 질문',.58,5.29,2.8,.42,18,True,PURPLE)
txt(s,'어떤 모델·부하에서 실행을 나누고,\n어떤 장비와 컴파일 조건을 선택할 것인가?',.58,5.91,12.1,.98,27,True)
nt(s,'앞의 세 작업을 정책 연구의 실험 기반으로 연결하는 제안이며 통합 자동 최적화기를 이미 완성했다는 의미가 아님. IR/컴파일 비용 모델, KV 이동을 고려한 단계 배치, 관측 기반 정책 평가를 협업 후보로 제시. 과제 총괄 배경은 다른 발표자가 작성하므로 생략.')
p.core_properties.title='ETRI — 저수준 NPU 실행 분석에서 이종 자원 관리까지'
p.save(OUT/'ETRI_저수준NPU연구와_자원관리_v7.pptx')
print('saved v7')
