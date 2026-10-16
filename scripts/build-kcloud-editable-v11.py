"""Native, individually editable reconstruction of the approved v10 diagrams."""
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from lxml import etree
import json, zipfile

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs/presentations/ETRI_NPU_Editable_v11'
OUT.mkdir(parents=True,exist_ok=True)
A=ROOT/'docs/presentations/kcloud-etri-research-v7/assets'
LOGOS=ROOT/'docs/presentations/kcloud-etri-visual-v6/assets'
REF=ROOT/'docs/presentations/ETRI_NPU_TopDown_v10/origin_image/slide_01.png'
p=Presentation();p.slide_width=Inches(16);p.slide_height=Inches(9)
BLUE='005EB8'; NAVY='09235D'; MUTED='5C6D84'; PALE='EAF4FD'; GRAY='EDF0F4'; WHITE='FFFFFF'
def txt(s,v,x,y,w,h,size=22,bold=False,c=NAVY,align=None):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));sh.name=v[:45]
    tf=sh.text_frame;tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(v.split('\n')):
        q=tf.paragraphs[0] if not i else tf.add_paragraph();q.text=line
        q.font.name='Noto Sans CJK KR';q.font.size=Pt(size);q.font.bold=bold;q.font.color.rgb=RGBColor.from_string(c)
        if align is not None:q.alignment=align
    return sh
def box(s,x,y,w,h,c=PALE,kind=MSO_SHAPE.ROUNDED_RECTANGLE,line=None):
    sh=s.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(c)
    if line:sh.line.color.rgb=RGBColor.from_string(line);sh.line.width=Pt(2)
    else:sh.line.fill.background()
    # Remove theme-inherited shadows, keeping all geometry native.
    sh._element.spPr.append(etree.Element('{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst'))
    for effect in sh._element.xpath('.//a:effectRef'):effect.set('idx','0')
    if kind==MSO_SHAPE.ROUNDED_RECTANGLE:sh.adjustments[0]=.10
    return sh
def line(s,x,y,x2,y2,c=BLUE,width=2):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x),Inches(y),Inches(x2),Inches(y2));sh.line.color.rgb=RGBColor.from_string(c);sh.line.width=Pt(width)
    for effect in sh._element.xpath('.//a:effectRef'):effect.set('idx','0')
    return sh
def arrow(s,x,y,w=.6,h=.35):return box(s,x,y,w,h,'84B8E8',MSO_SHAPE.RIGHT_ARROW)
def icon(s,kind,x,y,k=1,c=BLUE):
    # Icons are ordinary editable PowerPoint shapes, never rasterized.
    if kind=='model':
        pts=[(.12,.3),(.12,.75),(.5,.12),(.5,.5),(.5,.88),(.88,.3),(.88,.75)]
        for a,b in [(0,2),(0,3),(1,3),(1,4),(2,5),(3,5),(3,6),(4,6)]:line(s,x+pts[a][0]*k,y+pts[a][1]*k,x+pts[b][0]*k,y+pts[b][1]*k,'A7BAD0',1.4)
        for a,b in pts:box(s,x+(a-.08)*k,y+(b-.08)*k,.16*k,.16*k,c,MSO_SHAPE.OVAL)
    elif kind=='chip':
        box(s,x+.17*k,y+.17*k,.66*k,.66*k,PALE,MSO_SHAPE.ROUNDED_RECTANGLE,c)
        for q in [.3,.5,.7]:
            line(s,x+q*k,y,x+q*k,y+.17*k,c);line(s,x+q*k,y+.83*k,x+q*k,y+k,c)
            line(s,x,y+q*k,x+.17*k,y+q*k,c);line(s,x+.83*k,y+q*k,x+k,y+q*k,c)
    elif kind=='db':box(s,x+.1*k,y,.8*k,k,PALE,MSO_SHAPE.CAN,c)
    elif kind=='bars':
        for i,h in enumerate([.35,.6,.9]):box(s,x+i*.3*k,y+(1-h)*k,.2*k,h*k,c,MSO_SHAPE.RECTANGLE)
    else:
        box(s,x+.12*k,y,.72*k,k,WHITE,MSO_SHAPE.FOLDED_CORNER,c)
        if kind=='code':txt(s,'</>',x+.18*k,y+.32*k,.6*k,.4*k,18*k,True,c,PP_ALIGN.CENTER)
        else:
            for dy in [.3,.5,.7]:line(s,x+.27*k,y+dy*k,x+.65*k,y+dy*k,c,2)
def pic(s,path,x,y,w,h,crop=None):
    iw,ih=Image.open(path).size;l,t,r,b=crop or (0,0,iw,ih);f=min(w/(r-l),h/(b-t));ww,hh=(r-l)*f,(b-t)*f
    sh=s.shapes.add_picture(str(path),Inches(x+(w-ww)/2),Inches(y+(h-hh)/2),width=Inches(ww),height=Inches(hh))
    sh.crop_left=l/iw;sh.crop_top=t/ih;sh.crop_right=(iw-r)/iw;sh.crop_bottom=(ih-b)/ih
    return sh
def start(title,sub=''):
    s=p.slides.add_slide(p.slide_layouts[6]);txt(s,title,.5,.25,13.1,.9,43,True)
    pic(s,REF,13.8,.3,1.7,.65,(1443,30,1625,101))
    if sub:txt(s,sub,.53,1.14,13,.5,21,False,MUTED)
    return s
def center(s,v,x,y,w,h,size=24,bold=True,c=NAVY):return txt(s,v,x,y,w,h,size,bold,c,PP_ALIGN.CENTER)

s=start('이종 자원 관리, 우리는 어디를 개발했나','모델 운영에서 저수준 실행까지')
for y,label in [(1.9,'모델·서비스'),(3.26,'모델 운영'),(4.85,'컴파일·런타임'),(7.18,'GPU·NPU')]:txt(s,label,.55,y,2.2,.5,24,True)
box(s,3,1.85,9.45,1.02,GRAY);icon(s,'file',3.5,2.03,.58);center(s,'사용자 요청 → LLM 추론 서비스',4.5,2.13,7.2,.45,25)
box(s,3,3.13,9.45,1.05,BLUE);center(s,'lmd-top',3.5,3.25,8.4,.43,30,True,WHITE);center(s,'모델 준비 · 배포 · 관측',3.5,3.75,8.4,.35,20,False,WHITE)
for x,k,t in [(3,'file','저수준 실행 분석'),(7.84,'code','커스텀 실행 · KV 전송')]:box(s,x,4.46,4.61,1.02);icon(s,k,x+.2,4.68,.57);txt(s,t,x+.97,4.75,3.55,.45,22,True)
box(s,3,5.69,9.45,.58,GRAY);center(s,'벤더 컴파일러 · 실행 엔진',3.3,5.8,8.8,.35,22)
for y in [2.91,4.21,5.5,6.29]:box(s,7.53,y,.38,.2,'A7BAD0',MSO_SHAPE.UP_ARROW)
box(s,3,6.5,9.45,1.94,GRAY)
for x,name,crop in [(3.2,'NVIDIA GPU',(377,744,647,848)),(6.36,'Furiosa RNGD',(718,737,1002,849)),(9.51,'Rebellions ATOM',(1101,704,1194,854))]:pic(s,REF,x,6.61,2.8,1.25,crop);center(s,name,x,7.97,2.7,.34,19)
box(s,13,1.85,2.5,4.43,GRAY);center(s,'활용 기반',13.2,2.03,2.1,.4,23)
for y,file,title,sub in [(2.73,'kubernetes','Kubernetes','배포·실행 관리'),(4.48,'prometheus','Prometheus','지표 수집')]:pic(s,LOGOS/(file+'.png'),13.73,y,.94,.94);center(s,title,13.1,y+1,2.3,.35,18);center(s,sub,13.1,y+1.4,2.3,.3,14,False)
box(s,13,6.65,.28,.28,PALE);txt(s,'ETRI 분석·개발',13.43,6.6,2.1,.4,17,True,BLUE)
box(s,13,7.26,.28,.28,GRAY);txt(s,'기존 인프라',13.43,7.2,2,.4,17)
txt(s,'작업 범위의 개념도 · 단일 자동 제어기 통합을 의미하지 않음',3,8.64,12,.25,12,False,MUTED)

s=start('장비뿐 아니라 실행 조건도 선택해야 합니다','전체 환경 → 실행 조건')
box(s,.7,2.25,3.3,3.7);icon(s,'model',1.7,2.72,1.3);center(s,'모델',1,4.65,2.7,.5,30)
box(s,5.1,1.93,5.25,4.36);center(s,'조건별 아티팩트',5.35,2.14,4.7,.45,27)
for y,t in [(2.93,'입력 길이 설정'),(4,'배치 설정'),(5.07,'병렬성 설정')]:box(s,5.4,y,4.65,.9,WHITE);icon(s,'file',5.66,y+.15,.52);txt(s,t,6.53,y+.23,3.2,.4,23,True)
arrow(s,4.24,3.82);arrow(s,10.57,3.82)
box(s,11.46,2.25,3.7,3.7);pic(s,REF,12.12,2.66,2.35,2.25,(1101,704,1194,854));center(s,'NPU 실행',11.73,5.22,3.1,.45,28)
for x,k,title,desc in [(.7,'bars','실행 성능','같은 모델도 조건에 따라 달라짐'),(8.1,'file','컴파일 시간 · 파일 크기','실행 성능과 준비 비용을 함께 분석')]:box(s,x,6.8,7.05,1.42);icon(s,k,x+.25,7.09,.6);txt(s,title,x+1.12,7.02,5.6,.45,25,True);txt(s,desc,x+1.12,7.68,5.6,.35,17,False,MUTED)

s=start('모델이 NPU에서 실행되는 과정을 분석했습니다')
box(s,5.05,1.64,10.4,5.43);center(s,'우리가 분석한 범위',5.5,1.83,9.5,.5,28,True,BLUE)
for x,title,k,desc in [(.55,'모델','model','어떤 연산으로 구성되나?'),(5.6,'실행 준비','file','장비가 실행할 형태로 변환'),(11,'NPU 실행','chip','연산과 데이터 이동')]:
    box(s,x,2.7,4.2,2.95,GRAY if x<1 else WHITE);center(s,title,x+.2,2.91,3.8,.45,28);icon(s,k,x+1.4,3.61,1.3);center(s,desc,x+.15,5.08,3.9,.4,18,False,MUTED)
arrow(s,4.78,3.98,.7);arrow(s,10.03,3.98,.7)
txt(s,'저수준 IR',5.55,6.03,2.2,.44,24,True);txt(s,'연산·메모리 동작을 표현한 중간 단계',7.87,6.03,7.1,.45,20)
txt(s,'무엇을 계산하고, 데이터를 어디로 옮기는지 확인',5.55,6.57,9.3,.4,21,False,MUTED)
center(s,'실행 시간에 영향을 주는 지점과 변경 가능한 부분 파악',.8,7.72,14.5,.7,29,True,BLUE)

s=start('실행의 일부를 바꿔볼 수 있는 기반을 만들었습니다')
icon(s,'code',6.05,1.55,.8);txt(s,'우리가 개발한 연결 도구',7.15,1.76,6.6,.5,27,True,BLUE)
box(s,7.78,2.4,.4,.55,'A7BAD0',MSO_SHAPE.DOWN_ARROW)
box(s,1,3,14,2.63,GRAY,line='ACBACB')
for x,title,c in [(1.45,'기존 실행','DCE3EC'),(6.02,'직접 작성한 코드',BLUE),(10.62,'기존 실행','DCE3EC')]:box(s,x,3.45,3.87,1.24,c);center(s,title,x+.1,3.81,3.67,.48,28,True,WHITE if c==BLUE else NAVY)
arrow(s,5.44,3.86,.46);arrow(s,10.01,3.86,.46);center(s,'기존 NPU 실행 엔진',3,4.97,10,.5,26)
for x,k,v in [(1,'code','엔진 전체를 바꾸지 않고\n필요한 부분만 수정'),(8.3,'file','수정한 코드가 실제\n호출되는 경로 확인')]:box(s,x,6.06,6.7,1.55);icon(s,k,x+.3,6.4,.74);txt(s,v,x+1.33,6.34,5.05,1,24,True)
center(s,'실행 일부 변경 → 실행 단계를 다른 NPU로 분리',1,8.06,14,.6,27,True,BLUE)

s=start('실제 KV를 넘겨 NPU 사이에서 실행을 나눴습니다','커스텀 실행 → 단계 분리')
for x,title,desc in [(.5,'입력 처리 · Prefill','입력 토큰을 처리하여\n중간 상태(KV)를 생성'),(5.7,'중간 상태 · KV cache','생성된 KV를 저장하고\n다른 NPU로 전달'),(10.9,'답변 생성 · Decode','전달받은 KV로 이어서\n답변 토큰을 생성')]:box(s,x,1.95,4.6,4.79);center(s,title,x+.15,2.2,4.3,.48,25);center(s,desc,x+.2,2.95,4.2,.85,20,False,MUTED)
pic(s,REF,1.24,4.05,3.1,1.25,(718,737,1002,849));icon(s,'db',7.23,4.08,1.33);pic(s,REF,12.2,3.95,2,1.8,(1101,704,1194,854))
arrow(s,4.6,4.6,1.35);arrow(s,9.5,4.6,1.35)
center(s,'Furiosa RNGD',.8,5.85,4,.4,23);center(s,'직접 구현한 KV 전송 경로',5.84,5.84,4.3,.5,21,True,BLUE);center(s,'Rebellions ATOM',11.1,5.85,4.1,.4,23)
box(s,.5,7.14,15,1.05);txt(s,'KV 유지 · 연산 배치 · 컴파일 조건 분석',.82,7.47,8.3,.42,23,True);txt(s,'동종·이종 조합 / 역할 교환',9.6,7.47,5.4,.42,22,True,BLUE)

CAP=OUT/'assets'
with Image.open(CAP/'library.gif') as frame:frame.convert('RGB').save(CAP/'library.png')
for title,sub,points,asset in [
    ('lmd-top ① 모델·산출물 관리','모델을 찾고, 실행 가능한 빌드를 구분',
     [('원본과 빌드 구분','모델별 컴파일 산출물을\n계층으로 묶어 확인'),('벤더·실행 조건 확인','GPU·NPU 타깃과\n컴파일 설정을 함께 표시'),('배포 준비 상태 파악','바로 실행할 모델과\n컴파일이 필요한 모델 구분')],CAP/'library.png'),
    ('lmd-top ② 컴파일·배포 작업','준비한 모델을 실제 서빙으로 연결',
     [('벤더 컴파일 Job 관리','컴파일러 자체가 아니라\n컴파일 작업을 관리'),('작업 진행·실패 확인','Activity에서 진행률과\n실패·배포 상태 확인'),('배포 후 상태까지 연결','실행 중·시도 중·실패를\n구분하고 로그 확인')],A/'lmd-top-real.png'),
    ('lmd-top ③ 실행·자원 관측','배포 이후 모델과 장비 상태를 함께 확인',
     [('클러스터 상태 요약','모델 서빙과 가속기 현황을\n한 화면에서 확인'),('장비별 자원 관측','사용률·메모리·전력 등\n수집 가능한 지표 표시'),('기존 관측 기반 활용','Kubernetes·Prometheus 연계\n실제 화면 전환을 GIF로 표시')],CAP/'observe.gif')]:
    s=start(title,sub)
    for idx,(heading,desc) in enumerate(points):
        y=2.1+idx*1.87;box(s,.5,y,4.07,1.58);icon(s,['file','code','bars'][idx],.72,y+.2,.42)
        txt(s,heading,1.27,y+.18,3.1,.43,21,True,BLUE);txt(s,desc,.76,y+.77,3.55,.72,17,False,MUTED)
    if '②' in title:
        for x,label in [(5.05,'모델 선택'),(8.64,'컴파일·배포'),(12.2,'상태·로그 확인')]:
            box(s,x,2.25,3.12,1.15);center(s,label,x+.1,2.59,2.92,.45,22)
        arrow(s,8.25,2.66,.3);arrow(s,11.82,2.66,.3)
        txt(s,'Activity: 컴파일·배포 작업을 한 피드에서 확인',5.04,4.04,10.25,.5,24,True,BLUE)
        im=pic(s,asset,4.9,4.91,10.6,2.28,(10,586,1280,790))
        txt(s,'Running / Failed / Degraded / Serving 상태와 진행률 표시',5.06,7.68,10.1,.5,19,False,MUTED)
    else:im=pic(s,asset,4.88,1.86,10.62,6.74)
    assert im.image.blob==asset.read_bytes()
    txt(s,'실제 터미널 화면 · 캡처 당시 상태 · 외부 작업 실행 없이 관측 모드로 확보',4.98,8.68,10.6,.25,11,False,MUTED)

s=start('다음은 조건·단계·장비를 함께 고르는 정책입니다')
txt(s,'확보한 기반',.65,1.88,8,.5,27,True,BLUE)
for y,k,title,desc in [(2.64,'file','실행 조건','모델 특성 · 실행 환경 분석'),(4.34,'code','실행 단계','저수준 경로 분석 · 커스텀 실행'),(6.04,'chip','모델 운영','모델 준비 · 배포 · 관측')]:box(s,.6,y,8.7,1.38);icon(s,k,.98,y+.27,.75);txt(s,title,2.1,y+.18,6.7,.43,27,True);txt(s,desc,2.1,y+.84,6.7,.37,21,False,MUTED)
arrow(s,9.57,4.63,.65,.45);box(s,10.48,2.04,5,5.63,WHITE,line='7FABD6');center(s,'협업할 질문',10.7,2.29,4.55,.5,28,True,BLUE)
for y,t in [(3.42,'언제 실행을\n나눌까?'),(5.42,'어떤 설정·장비를\n고를까?')]:box(s,10.85,y,4.25,1.55);center(s,t,11,y+.26,3.95,1.1,27)
center(s,'비용 예측과 자원 배치 정책으로 연결',1,8.12,14,.6,29,True,BLUE)

out=OUT/'ETRI_NPU_Editable_v11.pptx';p.save(out)
report=[]
for i,s in enumerate(p.slides,1):report.append({'slide':i,'editable_textboxes':sum(sh.has_text_frame and bool(sh.text) for sh in s.shapes),'objects':len(s.shapes)})
with zipfile.ZipFile(out) as z:
    gifs=[n for n in z.namelist() if n.endswith('.gif')]
    assert len(gifs)==1 and z.read(gifs[0])==(CAP/'observe.gif').read_bytes()
print(json.dumps(report,ensure_ascii=False));print(out)
