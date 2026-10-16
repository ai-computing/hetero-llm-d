from pathlib import Path
import hashlib
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1]
DECK = ROOT / 'docs/presentations/ETRI_NPU_TopDown_v10'
ASSETS = ROOT / 'docs/presentations/kcloud-etri-renewal-v8/assets'
SCREEN = ROOT / 'docs/presentations/kcloud-etri-research-v7/assets/lmd-top-real.png'
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(16), Inches(9)
NAVY, BLUE, GRAY = '09235D', '005EB8', '596B84'

def text(slide, value, x, y, w, h, size=22, color=NAVY, bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for idx, line in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = 'Noto Sans CJK KR'
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
    return box

for number in range(1, 8):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    if number != 6:
        s.shapes.add_picture(str(DECK / f'origin_image/slide_{number:02d}.png'), 0, 0, width=prs.slide_width, height=prs.slide_height)
        continue
    text(s, 'lmd-top: 모델 준비부터 운영까지', .5, .3, 12.7, .7, 38, bold=True)
    text(s, '우리가 개발한 이종 GPU·NPU 모델 운영 도구', .53, 1.05, 12, .5, 23, GRAY)
    text(s, 'ETRI', 13.8, .25, 1.8, .6, 39, BLUE, True)
    text(s, '한국전자통신연구원', 13.85, .94, 1.85, .3, 10, BLUE)
    cards = [
        ('LuFolderOpen.png', '모델·산출물 관리', '원본 모델과 컴파일 결과 구분\n벤더별 모델·배포 가능 상태 확인'),
        ('LuRocket.png', '컴파일·배포 작업', '벤더 컴파일 Job과 배포 관리\n진행률·실패 상태·로그 확인'),
        ('LuActivity.png', '실행·자원 관측', '모델 서빙·요청·장비 상태 확인\nKubernetes·Prometheus 연계'),
    ]
    for idx, (icon, title, desc) in enumerate(cards):
        y = 2.0 + idx * 1.75
        shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.5), Inches(y), Inches(4.05), Inches(1.48))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string('ECF5FD')
        shape.line.fill.background()
        s.shapes.add_picture(str(ASSETS / icon), Inches(.7), Inches(y+.21), width=Inches(.4), height=Inches(.4))
        text(s, title, 1.23, y+.17, 3.1, .4, 21, BLUE, True)
        text(s, desc, .73, y+.72, 3.68, .65, 15, GRAY)
    text(s, '컴파일러 자체가 아니라\n모델 준비·운영 작업을 관리', .7, 7.45, 3.9, .8, 17, GRAY)
    picture = s.shapes.add_picture(str(SCREEN), Inches(4.9), Inches(1.85), width=Inches(10.55))
    assert picture.image.blob == SCREEN.read_bytes(), 'Original screenshot changed'
    text(s, '실제 컬러 터미널 화면 · 위: 모델/빌드 목록 · 아래: 작업 상태', 5.0, 8.67, 10.6, .25, 12, GRAY)

out = DECK / 'ETRI_NPU_TopDown_v10.pptx'
prs.save(out)
print(out)
print('7 slides; no speaker notes; original screenshot SHA256=' + hashlib.sha256(SCREEN.read_bytes()).hexdigest())
