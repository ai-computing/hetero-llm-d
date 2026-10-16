"""Template-derived five-slide overview; preserves source masters/layouts."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR
import subprocess

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs/presentations/kcloud-etri-visual-v5'
OUT.mkdir(parents=True,exist_ok=True)
TEMPLATE=Path('/home/gspark/.codex/plugins/cache/openai-curated-remote/openai-templates/0.1.1/skills/artifact-template-simple-light-mode/assets/reference.pptx')
INK='202027'; MUTED='72727D'; PURPLE='7560D5'; PALE='F2EFFB'; GRAY='F4F4F6'; GREEN='48816A'
assets=OUT/'assets'; assets.mkdir(exist_ok=True)
js="""const React=require('react'),R=require('react-dom/server'),I=require('react-icons/lu'),sharp=require('sharp');const path=process.argv[1];(async()=>{for(const name of ['LuCpu','LuServer','LuBoxes','LuFolderOpen','LuSlidersHorizontal','LuGauge','LuWorkflow','LuPackageCheck','LuRocket','LuActivity','LuNetwork','LuMemoryStick','LuArrowRight','LuDatabase']){let svg=R.renderToStaticMarkup(React.createElement(I[name],{size:320,color:'#7560D5',strokeWidth:1.35}));await sharp(Buffer.from(svg)).png().toFile(path+'/'+name+'.png')}})();"""
subprocess.run(['node','-e',js,str(assets)],check=True,env={**__import__('os').environ,'NODE_PATH':'/tmp/kcloud-ppt-tools/node_modules'})
p=Presentation(str(TEMPLATE))
# Structural work first: select existing cover, icon columns, timeline,
# image-led layout, and comparison layout. Do not modify the reference.
ids=list(p.slides._sldIdLst)
keep=[0,6,16,7,10]
for i,e in enumerate(ids):
    p.slides._sldIdLst.remove(e)
    if i not in keep:p.part.drop_rel(e.rId)
for i in keep:p.slides._sldIdLst.append(ids[i])

def txt(s,text,x,y,w,h,size=22,bold=False,color=INK):
    sh=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=sh.text_frame; tf.clear(); tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
    tf.word_wrap=True
    for i,line in enumerate(text.split('\n')):
        q=tf.paragraphs[0] if i==0 else tf.add_paragraph();q.text=line
        q.font.name='Noto Sans CJK KR';q.font.size=Pt(size);q.font.bold=bold;q.font.color.rgb=RGBColor.from_string(color)
        q.space_after=Pt(0)
    return sh
def box(s,x,y,w,h,fill=GRAY):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill);sh.line.fill.background();sh.adjustments[0]=.08;return sh
def icon(s,name,x,y,size=.75):s.shapes.add_picture(str(assets/(name+'.png')), Inches(x), Inches(y),width=Inches(size),height=Inches(size))
def line(s,x1,y1,x2,y2,color='C7C4D3'):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1),Inches(y1),Inches(x2),Inches(y2));sh.line.color.rgb=RGBColor.from_string(color);sh.line.width=Pt(1.6)
def arrow(s,x,y):icon(s,'LuArrowRight',x,y,.48)
def base(s,n,kicker,title):
    # User requested different sizing, placements and illustrations: replace
    # content slots while retaining actual template package and slide layouts.
    for sh in list(s.shapes):sh._element.getparent().remove(sh._element)
    txt(s,kicker.upper(),.58,.37,11,.3,11,True,MUTED)
    txt(s,title,.58,.94,12.2,1.35,36,True)
    txt(s,'ETRI  /  NPU · 자원 관리',.58,7.05,8,.22,9,False,MUTED)
    txt(s,f'{n:02d}',12.3,7.02,.45,.3,10,False,MUTED)
def notes(s,text):s.notes_slide.notes_text_frame.text=text

s=p.slides[0];base(s,1,'Infrastructure','GPU와 두 종류의 NPU를\n하나의 서빙 환경에 연결했습니다')
box(s,.58,2.78,3.0,3.74,PALE);icon(s,'LuServer',.89,3.08,1.0)
txt(s,'공통 서빙 환경',.89,4.33,2.55,.5,23,True)
txt(s,'요청 → 모델 서버\n저장소 → 모델 로드',.89,5.1,2.5,1,17,False,MUTED)
line(s,3.58,4.53,4.2,4.53);line(s,4.2,3.36,4.2,5.91)
for y,label,sub in [(2.73,'NVIDIA GPU','GPU 모델 서버'),(4.12,'Furiosa NPU','벤더 런타임 기반 서빙'),(5.51,'Rebellions NPU','벤더 런타임 기반 서빙')]:
    line(s,4.2,y+.57,4.64,y+.57);icon(s,'LuCpu',4.84,y+.12,.9)
    txt(s,label,6.02,y+.1,4,.44,24,True);txt(s,sub,6.02,y+.64,4,.4,16,False,MUTED)
icon(s,'LuDatabase',10.71,3.09,.95);txt(s,'공유 모델\n저장소',10.4,4.35,2.1,1.0,23,True)
notes(s,'구축 환경의 논리 구성. CPU/관리/모니터링 계층과 서버 이름은 생략. 장비 아이콘은 제품 외형이 아닌 개념 표현이다. 공통 API/Gateway는 각 모델 서버로 연결하며 lmd-top은 프록시가 아니다. 모든 모델이 현재 가동 중이라는 뜻이 아니다. 근거: docs/ENVIRONMENT-GUIDE.md, manifests/README-model-store.md, manifests/epp/README.md. 2026-09-08 조회에서는 모델 워커가 0 replicas였다.')

s=p.slides[1];base(s,2,'Characterization','NPU를 직접 구동하며\n준비 조건과 실행 특성을 살폈습니다')
for x,ic,title,body in [(.58,'LuSlidersHorizontal','컴파일 조건','입력 길이 · 실행 형태'),(4.87,'LuGauge','실행 특성','동시 요청 · 지연 · 전력'),(9.05,'LuMemoryStick','데이터 이동','CPU–NPU · KV 전달')]:
    icon(s,ic,x,2.86,1.15);txt(s,title,x,4.32,3.7,.52,26,True);txt(s,body,x,5.06,3.7,.5,19,False,MUTED)
box(s,.58,6.06,12.16,.62,PALE);txt(s,'GPU 경로와 비교해, NPU는 실행 전 준비 조건도 관리 대상이었습니다.',.83,6.19,11.7,.34,18,False,PURPLE)
notes(s,'비교는 사용한 소프트웨어 경로에 한정한다. GPU도 컴파일/최적화가 가능하며 NPU 전체의 본질적 우열로 일반화하지 않는다. Furiosa benchmark: bucket/shape, 입력·출력 길이, concurrency, power/tokJ 분석. 별도 PD 개발에서는 host placement/copy, resident KV, batching 등 실험. 성능 수치와 세부 비교 결과는 본문에서 의도적으로 제외. 근거: PD disaggregation_HISTORY.md §§3–7,12. 다른 서버에서 작성된 기록이며 원시 로그를 재검증한 결과는 아니다.')

s=p.slides[2];base(s,3,'Model lifecycle / lmd-top','모델 준비부터 배포까지\nlmd-top에서 관리하도록 구현했습니다')
for i,(ic,title,sub) in enumerate([('LuFolderOpen','모델 선택','원본 · 모델 카탈로그'),('LuWorkflow','컴파일 실행','벤더 도구의 작업 관리'),('LuPackageCheck','산출물 관리','원본과 구분해 저장'),('LuRocket','장비에 배포','대상 · 실행 옵션 선택')]):
    x=.58+3.13*i;icon(s,ic,x+.13,2.9,1.02);txt(s,title,x,4.25,2.9,.6,24,True);txt(s,sub,x,5.0,2.9,.9,17,False,MUTED)
    if i<3:arrow(s,x+2.35,3.18)
box(s,.58,6.05,12.16,.63,PALE);icon(s,'LuActivity',.83,6.17,.37);txt(s,'작업 진행 상태 · 성공 / 실패 · 로그를 함께 확인',1.41,6.17,10.8,.38,20,False,PURPLE)
notes(s,'lmd-top은 자체 컴파일러가 아니라 벤더별 compile/prefetch Job을 생성하고 관측하는 관리 도구. 원본과 컴파일 산출물은 공유 저장소에서 별도 관리한다. 모델/설정 지원 범위가 제한되며 GPU도 같은 컴파일을 반드시 거친다는 뜻은 아니다. 주요 UI: Library/Zoo, Compile, Activity, Deploy. 근거: lmd-top/README.ko.md, CHANGELOG.md v0.34–0.36, manifests/README-model-store.md.')

s=p.slides[3];base(s,4,'Operations / lmd-top','실행 중인 모델과 자원을\n한 화면에서 보고 운영합니다')
for y,ic,title,sub in [(2.95,'LuBoxes','모델 상태','배포 · 재시작 · 스케일'),(4.2,'LuNetwork','요청 경로','라우팅 · 선택 점수'),(5.45,'LuCpu','장비 상태','점유 · 메모리 · 전력')]:
    icon(s,ic,.6,y,.51);txt(s,title,1.3,y-.02,2.7,.45,23,True);txt(s,sub,1.3,y+.53,3.1,.45,16,False,MUTED)
s.shapes.add_picture(str(ROOT/'docs/images/lmd-top-ppt.gif'), Inches(4.55), Inches(2.54),width=Inches(8.19),height=Inches(4.87))
# Keep full screenshot at large scale, above the slide footer.
pic=s.shapes[-1];pic.top=Inches(2.35);pic.height=Inches(4.59);pic.width=Inches(7.71);pic.left=Inches(5.03)
txt(s,'UI 데모 · 수치 변화는 합성',9.2,7.04,2.8,.22,9,False,MUTED)
notes(s,'GIF는 실제 구성 기반 합성 UI 데모이며 실측 성능 결과가 아니다. PowerPoint 슬라이드 쇼에서 애니메이션 재생, PDF에서는 정지 화면. Serving/Traffic/Infra의 모델 상태·요청 경로·장비 상태를 보여준다. EPP 점수는 별도 score-observer 플러그인이 제공하는 경로에 한한다. 모든 route가 EPP를 통과하지 않는다. 운영 액션은 권한 모드와 확인 절차를 따른다. 근거: lmd-top/README.ko.md; lmd-top/contrib/epp-score-observer/README.md.')

s=p.slides[4];base(s,5,'Runtime exploration','한 모델의 실행 단계도\nNPU 사이에 나눠봤습니다')
box(s,.58,2.77,4.08,2.57);icon(s,'LuCpu',.88,3.1,.94);txt(s,'입력 문맥 처리',2.0,3.11,2.42,.5,23,True);txt(s,'Prefill · NPU A',2.0,3.81,2.4,.45,17,False,MUTED)
box(s,8.67,2.77,4.08,2.57);icon(s,'LuCpu',8.98,3.1,.94);txt(s,'답변 생성',10.09,3.11,2.42,.5,23,True);txt(s,'Decode · NPU B',10.09,3.81,2.4,.45,17,False,MUTED)
arrow(s,4.85,3.51);icon(s,'LuDatabase',6.19,3.03,.84);arrow(s,7.94,3.51)
txt(s,'중간 상태 전달',5.32,4.26,2.77,.47,21,True,PURPLE);txt(s,'KV cache',5.93,4.82,1.7,.36,15,False,MUTED)
txt(s,'Furiosa ↔ Rebellions',.6,5.83,4.6,.47,24,True)
txt(s,'동종·이종 조합과 역할 교환을 시험',.6,6.39,6.2,.44,19,False,MUTED)
txt(s,'협업 연결점',8.67,5.83,3.8,.4,16,True,PURPLE);txt(s,'어떤 모델·단계를 어디에 둘 것인가?',8.67,6.34,4.2,.6,20,True)
notes(s,'다른 서버에서 수행한 독립 런타임 연구. 자체 KV 입출력 경로에서 실제 KV를 전달하여 decode를 이어가는 PD 구현이며 초기 re-prefill 에뮬레이션과 구분한다. shared /dev/shm 기반 서빙 및 별도 TCP 테스트 기록. NPU A/B는 역할이고 서로 다른 물리 서버 간 전송을 의미하지 않는다. Furiosa/RNGD와 Rebellions/ATOM의 동종/이종 네 조합 및 역할 교환 실험. llm-d 배포 연계 이력은 있지만 lmd-top에 자동 PD 제어기가 통합됐다는 의미가 아니다. 성능은 부하/전송/배치 조건에 종속되므로 개선 수치를 주장하지 않는다. 앞의 환경, 분석, 관리 도구가 모델/단계 배치 정책 연구의 연결점이라는 논리로 마무리한다. 근거: PD disaggregation_HISTORY.md §§4,6,12.')

p.core_properties.title='ETRI — NPU와 모델·자원 관리: 구축·분석·구현'
p.core_properties.author='ETRI'
p.save(OUT/'ETRI_NPU_모델자원관리_5slides_v5.pptx')
print(OUT)
